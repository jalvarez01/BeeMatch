"""Extracción de texto de documentos Word (.docx)."""

import io

from backend.infrastructure.loaders.errores import DocumentoDanadoError, DocumentoSinTextoError


def extraer_texto_docx(contenido: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentoDanadoError(f"python-docx no disponible: {exc}")

    try:
        documento = Document(io.BytesIO(contenido))
        partes = [p.text.strip() for p in documento.paragraphs if p.text.strip()]

        # Muchas hojas de vida ponen la experiencia dentro de tablas.
        for tabla in documento.tables:
            for fila in tabla.rows:
                celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                if celdas:
                    partes.append(" | ".join(celdas))
    except Exception as exc:
        # Zip truncado, XML roto, partes faltantes: para el coordinador es lo mismo.
        raise DocumentoDanadoError(f"DOCX dañado o ilegible: {exc}")

    texto = "\n".join(partes).strip()
    if not texto:
        raise DocumentoSinTextoError("El documento no contiene texto.")
    return texto
