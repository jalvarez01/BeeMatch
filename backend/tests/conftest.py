"""Fábricas de documentos de prueba para HU-17. No se versionan CVs reales."""

import io
import shutil
import struct
import tempfile
from pathlib import Path

import pytest

FIXTURES = Path(__file__).parent / "fixtures"

TEXTO_CV = "Desarrollador backend con experiencia en Java Spring Boot y microservicios"


def crear_pdf(texto: str = TEXTO_CV) -> bytes:
    """PDF mínimo de una página con texto seleccionable, sin librerías extra."""
    flujo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(flujo) + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = []
    for numero, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += b"%d 0 obj\n" % numero + cuerpo + b"\nendobj\n"
    inicio_xref = len(salida)
    salida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    for pos in posiciones:
        salida += b"%010d 00000 n \n" % pos
    salida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objetos) + 1,
        inicio_xref,
    )
    return bytes(salida)


def cifrar_pdf(pdf: bytes, clave_usuario: str, algoritmo: str = "AES-256") -> bytes:
    from pypdf import PdfReader, PdfWriter

    escritor = PdfWriter()
    escritor.append(PdfReader(io.BytesIO(pdf)))
    escritor.encrypt(user_password=clave_usuario, owner_password="propietario", algorithm=algoritmo)
    salida = io.BytesIO()
    escritor.write(salida)
    return salida.getvalue()


def crear_docx(texto: str = TEXTO_CV, celda: str = "Banco Ejemplo | Core bancario") -> bytes:
    from docx import Document

    documento = Document()
    documento.add_paragraph(texto)
    tabla = documento.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text, tabla.rows[0].cells[1].text = celda.split(" | ")
    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


def cifrar_docx(docx: bytes, clave: str = "secreto") -> bytes:
    msoffcrypto = pytest.importorskip("msoffcrypto")
    salida = io.BytesIO()
    msoffcrypto.OfficeFile(io.BytesIO(docx)).encrypt(clave, salida)
    return salida.getvalue()


def doc_ejemplo() -> bytes:
    return (FIXTURES / "cv_ejemplo.doc").read_bytes()


def marcar_doc_como_cifrado(doc: bytes) -> bytes:
    """Enciende la bandera fEncrypted del encabezado FIB (MS-DOC) de un .doc real."""
    import olefile

    with tempfile.TemporaryDirectory() as carpeta:
        ruta = Path(carpeta) / "cifrado.doc"
        ruta.write_bytes(doc)
        ole = olefile.OleFileIO(str(ruta), write_mode=True)
        flujo = bytearray(ole.openstream("WordDocument").read())
        (banderas,) = struct.unpack_from("<H", flujo, 0x0A)
        struct.pack_into("<H", flujo, 0x0A, banderas | 0x0100)
        ole.write_stream("WordDocument", bytes(flujo))
        ole.close()
        return ruta.read_bytes()


requiere_antiword = pytest.mark.skipif(
    shutil.which("antiword") is None, reason="antiword no está instalado en esta máquina"
)
